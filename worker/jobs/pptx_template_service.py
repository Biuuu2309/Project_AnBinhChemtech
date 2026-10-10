"""Fill An Bình Chemtech PPTX V2 template by verified shape ids."""

from __future__ import annotations

import logging
import shutil
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from jobs.pptx_mapping import REQUIRED_SHAPES, ShapeWrite, build_shape_writes

logger = logging.getLogger("worker.pptx")


class MissingShapeError(ValueError):
    pass


def _iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _index_text_shapes(prs: Presentation) -> dict[int, Any]:
    found: dict[int, Any] = {}
    for slide in prs.slides:
        for shape in _iter_shapes(slide.shapes):
            if not getattr(shape, "has_text_frame", False):
                continue
            found[int(shape.shape_id)] = shape
    return found


def _set_shape_text(shape: Any, text: str) -> None:
    tf = shape.text_frame
    if not tf.paragraphs:
        tf.text = text
        return
    para = tf.paragraphs[0]
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.text = text
    for extra in tf.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def apply_shape_writes(prs: Presentation, writes: list[ShapeWrite]) -> None:
    by_id = _index_text_shapes(prs)
    missing = [sid for sid in REQUIRED_SHAPES if sid not in by_id]
    if missing:
        raise MissingShapeError(f"PPTX template missing required shape id(s): {missing}")

    for write in writes:
        shape = by_id.get(write.shape_id)
        if shape is None:
            raise MissingShapeError(f"PPTX template missing shape id={write.shape_id}")
        _set_shape_text(shape, write.text)


def fill_pptx_template(
    template_path: Path,
    output_path: Path,
    job: dict[str, Any],
    customer: dict[str, Any],
) -> Path:
    if not template_path.is_file():
        raise FileNotFoundError(f"PPTX template not found: {template_path}")

    writes = build_shape_writes(job, customer)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # Copy first so media/rels stay intact, then mutate
    shutil.copy2(template_path, output_path)
    prs = Presentation(str(output_path))
    apply_shape_writes(prs, writes)
    prs.save(str(output_path))

    logger.info(
        "pptx_filled quotation_id=%s shapes_written=%s output=%s",
        job.get("id"),
        len(writes),
        output_path.name,
    )
    return output_path


def read_shape_texts(path: Path) -> dict[int, str]:
    """Test helper: shape_id → concatenated text."""
    prs = Presentation(str(path))
    by_id = _index_text_shapes(prs)
    out: dict[int, str] = {}
    for sid, shape in by_id.items():
        out[sid] = "".join(
            (run.text or "") for p in shape.text_frame.paragraphs for run in p.runs
        ) or (shape.text_frame.text or "")
    return out
