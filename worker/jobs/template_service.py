from pathlib import Path

from docx import Document


def _replace_in_paragraph(paragraph, mapping: dict[str, str]) -> None:
    if not paragraph.runs:
        text = paragraph.text
        for key, value in mapping.items():
            text = text.replace(f"{{{{{key}}}}}", value)
        if text != paragraph.text:
            paragraph.text = text
        return

    full = "".join(run.text for run in paragraph.runs)
    updated = full
    for key, value in mapping.items():
        updated = updated.replace(f"{{{{{key}}}}}", value)
    if updated == full:
        return

    paragraph.runs[0].text = updated
    for run in paragraph.runs[1:]:
        run.text = ""


def _replace_all(doc: Document, mapping: dict[str, str]) -> None:
    for paragraph in doc.paragraphs:
        _replace_in_paragraph(paragraph, mapping)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_paragraph(paragraph, mapping)


def build_field_mapping(job: dict, customer: dict) -> dict[str, str]:
    items = job.get("items") or []
    lines = []
    for idx, item in enumerate(items, start=1):
        qty = item.get("quantity", 0)
        price = item.get("unit_price", 0)
        line_total = float(qty) * float(price)
        lines.append(
            f"{idx}. {item.get('product_name', '')} | "
            f"{item.get('specification', '')} | "
            f"SL: {qty} | Đơn giá: {price:,.0f} | "
            f"Thành tiền: {line_total:,.0f}"
        )

    first = items[0] if items else {}
    total = sum(float(i.get("quantity", 0)) * float(i.get("unit_price", 0)) for i in items)

    return {
        "quotation_number": job.get("id", ""),
        "customer_name": customer.get("name", ""),
        "customer_company": customer.get("company", ""),
        "customer_email": customer.get("email", ""),
        "customer_phone": customer.get("phone", ""),
        "customer_address": customer.get("address", ""),
        "product_name": first.get("product_name", ""),
        "specification": first.get("specification", ""),
        "quantity": str(first.get("quantity", "")),
        "unit_price": f"{float(first.get('unit_price', 0)):,.0f}",
        "items_block": "\n".join(lines) if lines else "(no items)",
        "total_amount": f"{total:,.0f}",
        "payment_terms": job.get("payment_terms", ""),
        "delivery_terms": job.get("delivery_terms", ""),
        "note": job.get("note") or "",
    }


def fill_template(template_path: Path, output_path: Path, mapping: dict[str, str]) -> Path:
    if not template_path.is_file():
        raise FileNotFoundError(f"Template not found: {template_path}")

    doc = Document(template_path)
    _replace_all(doc, mapping)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path
