from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_quotation_template(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    title = doc.add_heading("BẢNG BÁO GIÁ / QUOTATION", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("Công ty: An Bình Chemtech")
    doc.add_paragraph("Số báo giá: {{quotation_number}}")
    doc.add_paragraph("")

    doc.add_heading("Thông tin khách hàng", level=2)
    doc.add_paragraph("Người liên hệ: {{customer_name}}")
    doc.add_paragraph("Công ty: {{customer_company}}")
    doc.add_paragraph("Email: {{customer_email}}")
    doc.add_paragraph("Điện thoại: {{customer_phone}}")
    doc.add_paragraph("Địa chỉ: {{customer_address}}")
    doc.add_paragraph("")

    doc.add_heading("Chi tiết sản phẩm", level=2)
    doc.add_paragraph("Sản phẩm (dòng 1): {{product_name}}")
    doc.add_paragraph("Quy cách: {{specification}}")
    doc.add_paragraph("Số lượng: {{quantity}}")
    doc.add_paragraph("Đơn giá (VND): {{unit_price}}")
    doc.add_paragraph("")
    doc.add_paragraph("Danh sách mặt hàng:")
    doc.add_paragraph("{{items_block}}")
    doc.add_paragraph("")
    doc.add_paragraph("Tổng cộng (VND): {{total_amount}}")
    doc.add_paragraph("")

    doc.add_heading("Điều kiện", level=2)
    doc.add_paragraph("Thanh toán: {{payment_terms}}")
    doc.add_paragraph("Giao hàng: {{delivery_terms}}")
    doc.add_paragraph("Ghi chú: {{note}}")
    doc.add_paragraph("")
    doc.add_paragraph("Trân trọng,")
    doc.add_paragraph("An Bình Chemtech")

    doc.save(path)
    return path


if __name__ == "__main__":
    target = Path(__file__).resolve().parents[1] / "templates" / "quotation_template.docx"
    create_quotation_template(target)
    print(f"Created {target}")
